"""Content extraction: text, images, video frames, audio clips and thumbnails."""
import io
import json
import os
import re
import subprocess
import tempfile
from dataclasses import dataclass, field

import numpy as np
from PIL import Image, ImageOps

try:
    import pillow_heif
    pillow_heif.register_heif_opener()
except Exception:
    pass

Image.MAX_IMAGE_PIXELS = 400_000_000
SF_DATALESS = 0x40000000
THUMB_PX = 360
EMBED_PX = 896
MAX_TEXT = 24_000
CHUNK = 1800
MAX_CHUNKS = 4


@dataclass
class Extracted:
    text: str = ""
    images: list = field(default_factory=list)   # PIL images (visual vectors)
    audio: object = None                           # float32 mono 16 kHz
    thumb: object = None                           # PIL image
    meta: dict = field(default_factory=dict)


def is_dataless(st) -> bool:
    return bool(getattr(st, "st_flags", 0) & SF_DATALESS)


def _rgb(im: Image.Image) -> Image.Image:
    im = ImageOps.exif_transpose(im)
    if im.mode in ("RGBA", "LA", "P"):
        im = im.convert("RGBA")
        bg = Image.new("RGB", im.size, (255, 255, 255))
        bg.paste(im, mask=im.split()[-1])
        return bg
    return im.convert("RGB")


def _fit(im, px):
    im = im.copy()
    im.thumbnail((px, px), Image.LANCZOS)
    return im


def save_thumb(im: Image.Image, dest):
    t = _fit(im, THUMB_PX)
    t.save(dest, "WEBP", quality=72, method=4)


def quicklook(path, px=1024, timeout=25):
    """Render a thumbnail with macOS Quick Look (handles PSD, AI, RAW, SVG, Office, Pages...)."""
    with tempfile.TemporaryDirectory() as d:
        try:
            subprocess.run(["qlmanage", "-t", "-s", str(px), "-o", d, path],
                           capture_output=True, timeout=timeout)
        except subprocess.TimeoutExpired:
            return None
        for f in os.listdir(d):
            try:
                im = Image.open(os.path.join(d, f))
                im.load()
                return _rgb(im)
            except Exception:
                return None
    return None


def load_image(path, keep_alpha=False):
    im = Image.open(path)
    if im.format == "JPEG":
        im.draft("RGB", (EMBED_PX * 2, EMBED_PX * 2))
    if getattr(im, "is_animated", False):
        im.seek(0)
    im.load()
    if keep_alpha and (im.mode in ("RGBA", "LA") or (im.mode == "P" and "transparency" in im.info)):
        return ImageOps.exif_transpose(im).convert("RGBA")
    return _rgb(im)


def ffprobe(path):
    try:
        out = subprocess.run(
            ["ffprobe", "-v", "error", "-print_format", "json", "-show_format", "-show_streams", path],
            capture_output=True, timeout=30).stdout
        d = json.loads(out or b"{}")
    except Exception:
        return {}
    meta = {}
    try:
        meta["duration"] = float(d.get("format", {}).get("duration", 0) or 0)
    except ValueError:
        pass
    for s in d.get("streams", []):
        if s.get("codec_type") == "video" and s.get("width"):
            meta["width"], meta["height"] = s["width"], s["height"]
            break
    return meta


def video_frame(path, t, px=EMBED_PX):
    try:
        out = subprocess.run(
            ["ffmpeg", "-v", "error", "-ss", f"{t:.2f}", "-i", path, "-frames:v", "1",
             "-vf", f"scale='min({px},iw)':-2", "-f", "image2pipe", "-vcodec", "png", "-"],
            capture_output=True, timeout=60).stdout
        return _rgb(Image.open(io.BytesIO(out))) if out else None
    except Exception:
        return None


def audio_clip(path, start=0.0, secs=30):
    try:
        out = subprocess.run(
            ["ffmpeg", "-v", "error", "-ss", f"{start:.1f}", "-t", str(secs), "-i", path,
             "-ac", "1", "-ar", "16000", "-f", "f32le", "-"],
            capture_output=True, timeout=90).stdout
    except Exception:
        return None
    a = np.frombuffer(out, dtype=np.float32)
    return a.copy() if a.size > 16000 else None


def _read_text(path, limit=MAX_TEXT):
    with open(path, "rb") as f:
        raw = f.read(limit * 2)
    for enc in ("utf-8", "cp1250", "latin-1"):
        try:
            return raw.decode(enc)[:limit]
        except UnicodeDecodeError:
            continue
    return ""


def _textutil(path):
    try:
        out = subprocess.run(["textutil", "-convert", "txt", "-stdout", path],
                             capture_output=True, timeout=60).stdout
        return out.decode("utf-8", "ignore")[:MAX_TEXT]
    except Exception:
        return ""


def _pdf(path, ex: Extracted):
    import fitz
    doc = fitz.open(path)
    ex.meta["pages"] = doc.page_count
    parts = []
    for i, page in enumerate(doc):
        if i >= 10 or sum(map(len, parts)) > MAX_TEXT:
            break
        parts.append(page.get_text("text"))
    ex.text = "\n".join(parts)[:MAX_TEXT]
    if doc.page_count:
        p = doc[0]
        zoom = EMBED_PX / max(p.rect.width, p.rect.height, 1)
        pix = p.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False)
        im = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
        ex.thumb = im
        # Scans / graphic PDFs (little text) get a visual vector of page 1.
        if len(ex.text.strip()) < 200:
            ex.images.append(im)
    doc.close()


def _docx(path):
    import docx
    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables[:10]:
        for r in t.rows[:50]:
            parts.append(" | ".join(c.text for c in r.cells))
    return "\n".join(parts)[:MAX_TEXT]


def _pptx(path):
    from pptx import Presentation
    prs = Presentation(path)
    parts = []
    for i, s in enumerate(prs.slides):
        for sh in s.shapes:
            if getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip():
                parts.append(sh.text_frame.text)
        if sum(map(len, parts)) > MAX_TEXT:
            break
    return "\n".join(parts)[:MAX_TEXT]


def _xlsx(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets[:4]:
        parts.append(f"[{ws.title}]")
        for r in ws.iter_rows(max_row=150, values_only=True):
            vals = [str(v) for v in r if v is not None]
            if vals:
                parts.append(" | ".join(vals))
        if sum(map(len, parts)) > MAX_TEXT:
            break
    wb.close()
    return "\n".join(parts)[:MAX_TEXT]


def _html(path):
    t = _read_text(path, MAX_TEXT * 3)
    t = re.sub(r"(?is)<(script|style).*?</\1>", " ", t)
    t = re.sub(r"<[^>]+>", " ", t)
    return re.sub(r"\s+", " ", t)[:MAX_TEXT]


def _ipynb(path):
    try:
        nb = json.loads(_read_text(path, 2_000_000))
        return "\n".join("".join(c.get("source", [])) for c in nb.get("cells", []))[:MAX_TEXT]
    except Exception:
        return ""


def extract(path: str, kind: str, size: int) -> Extracted:
    ex = Extracted()
    ext = os.path.splitext(path)[1].lower()

    if kind == "image" and ext not in (".svg", ".ico"):
        if size > 300_000_000:
            return ex
        im = load_image(path, keep_alpha=True)
        ex.meta["width"], ex.meta["height"] = im.size
        if max(im.size) >= 96:  # icons / tiny responsive variants: thumbnail only
            ex.images.append(_rgb(_fit(im, EMBED_PX)))
        ex.thumb = im  # alpha kept so transparent logos show on a checkerboard
    elif kind in ("raw", "design", "font", "model3d") or ext in (".svg", ".ico", ".pages", ".key", ".numbers"):
        im = quicklook(path)
        if im is not None:
            ex.thumb = im
            if kind != "font":
                ex.images.append(_fit(im, EMBED_PX))
    elif kind == "video":
        m = ffprobe(path)
        ex.meta.update(m)
        dur = m.get("duration") or 0
        stamps = [dur * f for f in (0.08, 0.35, 0.62, 0.9)] if dur > 4 else [0.0]
        for t in stamps:
            fr = video_frame(path, t)
            if fr is not None:
                ex.images.append(fr)
        if ex.images:
            ex.thumb = ex.images[0]
    elif kind == "audio":
        m = ffprobe(path)
        ex.meta["duration"] = m.get("duration")
        dur = m.get("duration") or 0
        ex.audio = audio_clip(path, start=min(30.0, dur * 0.25) if dur > 90 else 0.0)
    elif kind == "pdf":
        _pdf(path, ex)
    elif ext == ".docx":
        ex.text = _docx(path)
    elif ext == ".pptx":
        ex.text = _pptx(path)
    elif ext in (".xlsx", ".xlsm"):
        ex.text = _xlsx(path)
    elif ext in (".doc", ".rtf", ".odt"):
        ex.text = _textutil(path)
    elif ext == ".ipynb":
        ex.text = _ipynb(path)
    elif kind == "web":
        ex.text = _html(path)
    elif kind in ("text", "code") or ext in (".csv", ".tsv"):
        if size < 20_000_000:
            ex.text = _read_text(path)
    ex.text = (ex.text or "").replace("\x00", " ")
    # Never keep full-resolution bitmaps around (a 1440x6703 screenshot is ~38 MB in RAM).
    if ex.thumb is not None:
        ex.thumb = _fit(ex.thumb, THUMB_PX)
    return ex


def quick_thumb(path: str, kind: str):
    """Cheap thumbnail without the model — used by the fast thumbnail pass and the server."""
    if kind == "image" and not path.lower().endswith((".svg", ".ico")):
        return _fit(load_image(path, keep_alpha=True), THUMB_PX)
    if kind == "video":
        dur = ffprobe(path).get("duration") or 0
        return video_frame(path, dur * 0.08 if dur > 4 else 0.0, px=THUMB_PX * 2)
    if kind == "pdf":
        import fitz
        doc = fitz.open(path)
        try:
            p = doc[0]
            zoom = THUMB_PX * 1.5 / max(p.rect.width, p.rect.height, 1)
            pix = p.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False)
            return Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
        finally:
            doc.close()
    return None


def chunks(text: str):
    import unicodedata
    text = unicodedata.normalize("NFC", text)
    text = re.sub(r"[ \t\r\f\v]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n", text).strip()
    if len(text) < 20:
        return []
    out = []
    i = 0
    while i < len(text) and len(out) < MAX_CHUNKS:
        j = min(len(text), i + CHUNK)
        if j < len(text):
            k = text.rfind(" ", i + CHUNK // 2, j)
            j = k if k > 0 else j
        out.append(text[i:j])
        i = j
    return out


def snippet(text: str, n=320):
    return re.sub(r"\s+", " ", text or "").strip()[:n]
